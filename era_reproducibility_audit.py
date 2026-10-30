#!/usr/bin/env python3
"""Reproduce the ERA workbook arithmetic and weight-sensitivity audit.

This script audits only the supplied synthetic workbook. It does not test
human agreement, model accuracy, fairness, or real-world validity.
"""
import argparse, csv, json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.stats import spearmanr

SEED = 20261027
WEIGHTS = np.array([0.25, 0.25, 0.25, 0.10, 0.15])  # EC, ER, CA, (1-CP), ET
ALT = {
    "equal": np.array([.20, .20, .20, .20, .20]),
    "evidence_heavy": np.array([.30, .30, .15, .10, .15]),
    "traceability_heavy": np.array([.20, .20, .20, .15, .25]),
}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", help="QC-corrected synthetic workbook")
    ap.add_argument("--outdir", default="era_audit_outputs")
    args = ap.parse_args()
    out = Path(args.outdir); out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(args.workbook, data_only=True, read_only=True)
    if "100_Cases" not in wb.sheetnames or "Annotation_Form" not in wb.sheetnames:
        raise SystemExit("Expected 100_Cases and Annotation_Form sheets")
    ws = wb["100_Cases"]
    data = list(ws.values); headers = list(data[0]); ix = {h:i for i,h in enumerate(headers)}
    required = ["case_id", "role", "case_pattern", "EC", "ER", "CA", "CP", "ET", "illustrative_ERA", "independent_benchmark_label"]
    missing = [x for x in required if x not in ix]
    if missing: raise SystemExit(f"Missing columns: {missing}")
    rows = data[1:]
    vals = np.array([[float(r[ix[c]]) for c in ["EC","ER","CA","CP","ET","illustrative_ERA"]] for r in rows])
    EC, ER, CA, CP, ET, stored = vals.T
    calculated = .25*EC + .25*ER + .25*CA + .15*ET + .10*(1-CP)
    diffs = np.abs(stored-calculated)
    ann = list(wb["Annotation_Form"].values)
    ah = {h:i for i,h in enumerate(ann[0])}
    core = ["evidence_relevant_1to5", "evidence_validity_1to5", "benchmark_criterion_rating_1to5", "contradiction_present_yes_no"]
    incomplete_rows = sum(1 for r in ann[1:] if any(r[ah[c]] is None or str(r[ah[c]]).strip()=="" for c in core))
    ids = [r[ix["case_id"]] for r in rows]
    roles = {}
    for r in rows: roles[str(r[ix["role"]])] = roles.get(str(r[ix["role"]]),0)+1
    patterns = {}
    for r in rows: patterns[str(r[ix["case_pattern"]])] = patterns.get(str(r[ix["case_pattern"]]),0)+1
    provisional = sum("requires human annotation" in str(r[ix["annotation_status"]]).lower() for r in rows)

    def score(w): return w[0]*EC + w[1]*ER + w[2]*CA + w[3]*(1-CP) + w[4]*ET
    base = score(WEIGHTS)
    named = {}
    for name, w in ALT.items():
        sc = score(w)
        named[name] = {"mean_composite": float(sc.mean()), "spearman_rho_vs_proposed": float(spearmanr(base,sc).statistic)}
    rng = np.random.default_rng(SEED)
    sensitivity = {}
    for name, alpha in [("local_dirichlet_40x_proposed", 40*WEIGHTS), ("broad_dirichlet_ones", np.ones(5))]:
        draws = rng.dirichlet(alpha, size=10000)
        rho = np.array([spearmanr(base, score(w)).statistic for w in draws])
        sensitivity[name] = {"draws":10000,"p05":float(np.quantile(rho,.05)),"median":float(np.median(rho)),"p95":float(np.quantile(rho,.95)),"share_rho_ge_0_90":float(np.mean(rho>=.90))}
        with (out/f"{name}.csv").open("w",newline="") as f:
            wr=csv.writer(f); wr.writerow(["draw","spearman_rho"]); wr.writerows((i+1,float(x)) for i,x in enumerate(rho))
    summary = {
      "scope":"synthetic workbook arithmetic and mathematical weight sensitivity only",
      "seed":SEED,"cases":len(rows),"unique_case_ids":len(set(ids)),"roles":roles,"case_patterns":patterns,
      "annotation_rows":len(ann)-1,"annotation_rows_with_any_core_field_missing":incomplete_rows,
      "provisional_benchmark_cases":provisional,
      "stored_era_mean":float(stored.mean()),"stored_era_min":float(stored.min()),"stored_era_max":float(stored.max()),
      "recomputed_era_mean":float(calculated.mean()),"mean_absolute_difference":float(diffs.mean()),
      "maximum_absolute_difference":float(diffs.max()),"rows_difference_gt_0_001":int(np.sum(diffs>.001)),
      "named_weight_schemes":named,"monte_carlo_sensitivity":sensitivity,
      "interpretation":"These results do not establish construct validity, criterion validity, AI accuracy, human agreement, fairness, explanation usefulness, or real-world effectiveness."
    }
    (out/"ERA_reproducibility_audit_summary.json").write_text(json.dumps(summary,indent=2)+"\n")
    with (out/"ERA_case_level_recomputed_scores.csv").open("w",newline="") as f:
        wr=csv.writer(f); wr.writerow(["case_id","stored_ERA","recomputed_ERA","absolute_difference"])
        wr.writerows((r[ix["case_id"]],float(s),float(c),float(d)) for r,s,c,d in zip(rows,stored,calculated,diffs))
    print(json.dumps(summary,indent=2))

if __name__ == "__main__": main()
